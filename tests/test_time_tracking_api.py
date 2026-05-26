import io
from datetime import date, time
from unittest.mock import Mock, patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import RefreshToken

from core.enums import (
    ImportBatchSource,
    ImportRowStatus,
    ProjectAssignmentStatus,
    ProjectStatus,
    ProjectType,
    TimeEntryAuditEventType,
    TimeEntrySourceChangeFlag,
    TimeEntrySourceType,
    TimeEntryStatus,
)
from core.models import (
    JiraConnection,
    JiraIssueMapping,
    JiraProjectMapping,
    JiraUserMapping,
    Permission,
    Project,
    ProjectAssignment,
    Role,
    TempoAccountMapping,
    TempoConnection,
    TempoProjectMapping,
    TempoTeamMapping,
    TempoUserMapping,
    TimeEntry,
    TimeEntryAuditEvent,
    TimeImportBatch,
    TimeTask,
)


class TimeTrackingAPITestCase(APITestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        from django.core.management import call_command

        call_command("setup_public_tenant", "--domain", "testserver", verbosity=0)

    def setUp(self):
        self.employee_role = Role.objects.create(name="employee")
        self.manager_role = Role.objects.create(name="manager")
        self.no_access_role = Role.objects.create(name="no_access")

        self.view_own = self._permission("view_own_timesheet")
        self.view_team = self._permission("view_team_timesheets")
        self.approve_team = self._permission("approve_team_timesheets")
        self.view_dept = self._permission("view_dept_timesheets")
        self.export_timesheets = self._permission("export_timesheets")

        self.employee_role.permissions.add(self.view_own)
        self.manager_role.permissions.add(
            self.view_own,
            self.view_team,
            self.approve_team,
            self.view_dept,
            self.export_timesheets,
        )

        self.manager = self._make_user("manager", self.manager_role)
        self.employee = self._make_user("employee", self.employee_role)
        self.outsider = self._make_user("outsider", self.no_access_role)
        self.employee.profile.managers.add(self.manager.profile)

        self.project = Project.objects.create(
            name="Alpha",
            client="Acme",
            project_type=ProjectType.CLIENT,
            status=ProjectStatus.ACTIVE,
            owner=self.manager.profile,
        )
        self.other_project = Project.objects.create(
            name="Beta",
            client="Globex",
            project_type=ProjectType.CLIENT,
            status=ProjectStatus.ACTIVE,
        )
        self.task = TimeTask.objects.create(
            project=self.project,
            name="Implementation",
            jira_issue_key="ALPHA-1",
            jira_project_key="ALPHA",
        )
        self.other_task = TimeTask.objects.create(
            project=self.other_project,
            name="Discovery",
        )

    def _permission(self, action):
        permission, _ = Permission.objects.get_or_create(
            module_name="Time Tracking",
            feature_action=action,
        )
        return permission

    def _make_user(self, username, role):
        user = User.objects.create_user(
            username=username,
            email=f"{username}@example.com",
            password="pass",
        )
        profile = user.profile
        profile.role = role
        profile.full_name = username.title()
        profile.save(update_fields=["role", "full_name"])
        return user

    def _auth(self, user):
        refresh = RefreshToken.for_user(user)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {refresh.access_token}")

    def _entry_payload(self, **overrides):
        payload = {
            "project_id": self.project.id,
            "task_id": self.task.id,
            "work_date": "2026-05-18",
            "start_time": "10:00:00",
            "hours": "2.50",
            "notes": "Build time tracking foundation",
        }
        payload.update(overrides)
        return payload

    def _time_entry(self, **overrides):
        payload = {
            "employee": self.employee.profile,
            "project": self.project,
            "task": self.task,
            "work_date": date(2026, 5, 18),
            "start_time": time(10, 0),
            "hours": "1.00",
            "notes": "Entry",
            "source_type": TimeEntrySourceType.MANUAL,
            "source_external_id": "",
            "source_metadata": {},
            "status": TimeEntryStatus.DRAFT,
            "duplicate_fingerprint": "",
        }
        payload.update(overrides)
        entry = TimeEntry.objects.create(**payload)
        from core.services.time_tracking_service import fingerprint_for_entry

        entry.duplicate_fingerprint = fingerprint_for_entry(entry)
        entry.save(update_fields=["duplicate_fingerprint"])
        return entry

    def test_employee_can_create_own_manual_entry(self):
        self._auth(self.employee)
        response = self.client.post(
            reverse("core:time-entry-list"),
            self._entry_payload(),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["employee_id"], self.employee.profile.id)
        self.assertEqual(response.data["source_type"], TimeEntrySourceType.MANUAL)
        self.assertEqual(response.data["start_time"], "10:00:00")
        self.assertEqual(response.data["status"], TimeEntryStatus.DRAFT)
        entry = TimeEntry.objects.get(pk=response.data["id"])
        self.assertEqual(entry.start_time, time(10, 0))
        self.assertTrue(entry.duplicate_fingerprint)
        self.assertEqual(
            entry.audit_events.get().event_type,
            TimeEntryAuditEventType.CREATED,
        )

    def test_manual_entry_preserves_iso_start_time_wall_clock(self):
        self._auth(self.employee)
        response = self.client.post(
            reverse("core:time-entry-list"),
            self._entry_payload(start_time="2026-05-18T10:00:00.000+02:00"),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["start_time"], "10:00:00")
        entry = TimeEntry.objects.get(pk=response.data["id"])
        self.assertEqual(entry.start_time, time(10, 0))

    def test_manual_entry_accepts_twelve_hour_start_time(self):
        self._auth(self.employee)
        response = self.client.post(
            reverse("core:time-entry-list"),
            self._entry_payload(start_time="12:00AM"),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["start_time"], "00:00:00")
        entry = TimeEntry.objects.get(pk=response.data["id"])
        self.assertEqual(entry.start_time, time(0, 0))

    def test_employee_cannot_create_weekend_entry(self):
        self._auth(self.employee)
        response = self.client.post(
            reverse("core:time-entry-list"),
            self._entry_payload(work_date="2026-05-23"),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(
            response.data["work_date"][0],
            "Time entries cannot be submitted for weekends.",
        )

    def test_employee_cannot_move_entry_to_weekend(self):
        entry = self._time_entry()
        self._auth(self.employee)
        response = self.client.patch(
            reverse("core:time-entry-detail", args=[entry.id]),
            {"work_date": "2026-05-24"},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(
            response.data["work_date"][0],
            "Time entries cannot be submitted for weekends.",
        )

    def test_user_without_time_permission_cannot_create_entry(self):
        self._auth(self.outsider)
        response = self.client.post(
            reverse("core:time-entry-list"),
            self._entry_payload(project_id=self.project.id, task_id=self.task.id),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_duplicate_fingerprint_is_reported(self):
        self._auth(self.employee)
        first = self.client.post(
            reverse("core:time-entry-list"),
            self._entry_payload(),
            format="json",
        )
        second = self.client.post(
            reverse("core:time-entry-list"),
            self._entry_payload(),
            format="json",
        )

        self.assertEqual(first.status_code, status.HTTP_201_CREATED)
        self.assertEqual(second.status_code, status.HTTP_201_CREATED)
        self.assertEqual(second.data["duplicate_of"], first.data["id"])

    def test_submit_week_and_manager_approval_flow(self):
        self._auth(self.employee)
        create_response = self.client.post(
            reverse("core:time-entry-list"),
            self._entry_payload(),
            format="json",
        )
        self.assertEqual(create_response.status_code, status.HTTP_201_CREATED)

        submit_response = self.client.post(
            reverse("core:time-entry-submit-week"),
            {"week_start": "2026-05-18"},
            format="json",
        )
        self.assertEqual(submit_response.status_code, status.HTTP_200_OK)
        self.assertEqual(submit_response.data[0]["status"], TimeEntryStatus.SUBMITTED)

        self._auth(self.manager)
        approve_response = self.client.post(
            reverse("core:time-entry-approve", args=[create_response.data["id"]]),
            format="json",
        )
        self.assertEqual(approve_response.status_code, status.HTTP_200_OK)
        self.assertEqual(approve_response.data["status"], TimeEntryStatus.APPROVED)

    def test_submit_week_skips_existing_weekend_entries(self):
        weekday_entry = self._time_entry(work_date=date(2026, 5, 18))
        weekend_entry = self._time_entry(work_date=date(2026, 5, 23))
        self._auth(self.employee)
        response = self.client.post(
            reverse("core:time-entry-submit-week"),
            {"week_start": "2026-05-18"},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual([row["id"] for row in response.data], [weekday_entry.id])
        weekday_entry.refresh_from_db()
        weekend_entry.refresh_from_db()
        self.assertEqual(weekday_entry.status, TimeEntryStatus.SUBMITTED)
        self.assertEqual(weekend_entry.status, TimeEntryStatus.DRAFT)

    def test_task_must_belong_to_selected_project(self):
        self._auth(self.employee)
        response = self.client.post(
            reverse("core:time-entry-list"),
            self._entry_payload(task_id=self.other_task.id),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("task_id", response.data)

    def test_imported_source_fields_are_immutable(self):
        self._auth(self.manager)
        response = self.client.post(
            reverse("core:time-entry-list"),
            self._entry_payload(
                employee_id=self.employee.profile.id,
                source_type=TimeEntrySourceType.JIRA,
                source_external_id="10001",
                source_metadata={
                    "issue_key": "ALPHA-1",
                    "worklog_id": "10001",
                    "started": "2026-05-18T09:00:00.000+0000",
                },
            ),
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)

        update_response = self.client.patch(
            reverse("core:time-entry-detail", args=[response.data["id"]]),
            {"source_external_id": "changed"},
            format="json",
        )
        self.assertEqual(update_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("source_external_id", update_response.data)

    def test_approver_can_delete_approved_imported_entry(self):
        entry = TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.project,
            task=self.task,
            work_date=date(2026, 5, 18),
            hours="1.00",
            notes="Imported approved",
            status=TimeEntryStatus.APPROVED,
            source_type=TimeEntrySourceType.TEMPO,
            source_external_id="tempo-delete",
            source_metadata={"tempo_worklog_id": "tempo-delete"},
            duplicate_fingerprint="tempo-delete",
        )
        self._auth(self.manager)

        response = self.client.delete(
            reverse("core:time-entry-detail", args=[entry.id])
        )

        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(TimeEntry.objects.filter(pk=entry.id).exists())

    def test_employee_cannot_delete_imported_entry(self):
        entry = TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.project,
            task=self.task,
            work_date=date(2026, 5, 18),
            hours="1.00",
            notes="Imported own",
            status=TimeEntryStatus.DRAFT,
            source_type=TimeEntrySourceType.TEMPO,
            source_external_id="tempo-own-delete",
            source_metadata={"tempo_worklog_id": "tempo-own-delete"},
            duplicate_fingerprint="tempo-own-delete",
        )
        self._auth(self.employee)

        response = self.client.delete(
            reverse("core:time-entry-detail", args=[entry.id])
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertTrue(TimeEntry.objects.filter(pk=entry.id).exists())

    def test_approved_manual_entry_still_cannot_be_deleted(self):
        entry = TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.project,
            task=self.task,
            work_date=date(2026, 5, 18),
            hours="1.00",
            notes="Manual approved",
            status=TimeEntryStatus.APPROVED,
            source_type=TimeEntrySourceType.MANUAL,
            duplicate_fingerprint="manual-approved-delete",
        )
        self._auth(self.manager)

        response = self.client.delete(
            reverse("core:time-entry-detail", args=[entry.id])
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertTrue(TimeEntry.objects.filter(pk=entry.id).exists())

    def test_reject_requires_reason(self):
        entry = TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.project,
            task=self.task,
            work_date=date(2026, 5, 18),
            hours="1.00",
            notes="Ready for review",
            status=TimeEntryStatus.SUBMITTED,
            duplicate_fingerprint="abc",
        )
        self._auth(self.manager)

        response = self.client.post(
            reverse("core:time-entry-reject", args=[entry.id]),
            {"reason": ""},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_weekly_summary_respects_midweek_allocation_changes(self):
        ProjectAssignment.objects.create(
            user_profile=self.employee.profile,
            project=self.project,
            allocation_percentage=50,
            start_date=date(2026, 5, 18),
            end_date=date(2026, 5, 20),
            status=ProjectAssignmentStatus.COMPLETED,
        )
        ProjectAssignment.objects.create(
            user_profile=self.employee.profile,
            project=self.other_project,
            allocation_percentage=25,
            start_date=date(2026, 5, 21),
            status=ProjectAssignmentStatus.ACTIVE,
        )
        TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.project,
            task=self.task,
            work_date=date(2026, 5, 18),
            hours="5.00",
            notes="Alpha work",
            duplicate_fingerprint="weekly-alpha",
        )
        TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.other_project,
            task=self.other_task,
            work_date=date(2026, 5, 21),
            hours="2.00",
            notes="Beta work",
            duplicate_fingerprint="weekly-beta",
        )
        self._auth(self.employee)

        response = self.client.get(
            reverse("core:time_tracking_weekly_summary"),
            {"week_start": "2026-05-18"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["planned_hours"], "16.00")
        self.assertEqual(response.data["actual_hours"], "7.00")
        self.assertEqual(response.data["unallocated_capacity_hours"], "24.00")
        projects = {row["project_name"]: row for row in response.data["projects"]}
        self.assertEqual(projects["Alpha"]["planned_hours"], "12.00")
        self.assertEqual(projects["Alpha"]["allocation_percentage"], "30.00")
        self.assertEqual(projects["Beta"]["planned_hours"], "4.00")
        self.assertEqual(projects["Beta"]["allocation_percentage"], "10.00")

    def test_weekly_summary_marks_actuals_without_allocation_unallocated(self):
        TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.other_project,
            task=self.other_task,
            work_date=date(2026, 5, 18),
            hours="3.00",
            notes="Unallocated work",
            duplicate_fingerprint="unallocated-beta",
        )
        self._auth(self.employee)

        response = self.client.get(
            reverse("core:time_tracking_weekly_summary"),
            {"week_start": "2026-05-18"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["planned_hours"], "0.00")
        self.assertEqual(response.data["actual_hours"], "3.00")
        self.assertEqual(
            response.data["projects"][0]["allocation_status"], "unallocated"
        )

    def test_weekly_summary_accepts_employee_alias(self):
        ProjectAssignment.objects.create(
            user_profile=self.employee.profile,
            project=self.project,
            allocation_percentage=100,
            start_date=date(2026, 5, 18),
            status=ProjectAssignmentStatus.ACTIVE,
        )
        self._auth(self.manager)

        response = self.client.get(
            reverse("core:time_tracking_weekly_summary"),
            {"week_start": "2026-05-18", "employee": self.employee.profile.id},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["employee_id"], self.employee.profile.id)
        self.assertEqual(response.data["planned_hours"], "40.00")

    def test_active_allocations_hide_ended_and_archived_assignments(self):
        archived = Project.objects.create(
            name="Archived",
            project_type=ProjectType.INTERNAL,
            status=ProjectStatus.ARCHIVED,
        )
        ProjectAssignment.objects.create(
            user_profile=self.employee.profile,
            project=self.project,
            allocation_percentage=60,
            start_date=date(2026, 5, 1),
            status=ProjectAssignmentStatus.ACTIVE,
        )
        ProjectAssignment.objects.create(
            user_profile=self.employee.profile,
            project=self.other_project,
            allocation_percentage=20,
            start_date=date(2026, 4, 1),
            end_date=date(2026, 5, 1),
            status=ProjectAssignmentStatus.COMPLETED,
        )
        ProjectAssignment.objects.create(
            user_profile=self.employee.profile,
            project=archived,
            allocation_percentage=20,
            start_date=date(2026, 5, 1),
            status=ProjectAssignmentStatus.ACTIVE,
        )
        self._auth(self.employee)

        response = self.client.get(
            reverse("core:time_tracking_active_allocations"),
            {"work_date": "2026-05-18"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["assignments"]), 1)
        self.assertEqual(response.data["assignments"][0]["project_id"], self.project.id)
        self.assertEqual(response.data["remaining_allocation_percentage"], "40.00")

    def test_manager_can_view_direct_report_weekly_summary(self):
        ProjectAssignment.objects.create(
            user_profile=self.employee.profile,
            project=self.project,
            allocation_percentage=100,
            start_date=date(2026, 5, 18),
            status=ProjectAssignmentStatus.ACTIVE,
        )
        self._auth(self.manager)

        response = self.client.get(
            reverse("core:time_tracking_weekly_summary"),
            {
                "employee_id": self.employee.profile.id,
                "week_start": "2026-05-18",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["planned_hours"], "40.00")

    def test_employee_cannot_view_other_employee_weekly_summary(self):
        self._auth(self.outsider)

        response = self.client.get(
            reverse("core:time_tracking_weekly_summary"),
            {
                "employee_id": self.employee.profile.id,
                "week_start": "2026-05-18",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def _jira_worklog(self, **overrides):
        worklog = {
            "id": "10001",
            "issueKey": "ALPHA-1",
            "issueId": "50001",
            "author": {
                "accountId": "acct-1",
                "displayName": "Employee Jira",
            },
            "started": "2026-05-18T09:00:00.000+0000",
            "timeSpentSeconds": 7200,
            "comment": "Jira implementation",
            "updated": "2026-05-18T11:00:00.000+0000",
        }
        worklog.update(overrides)
        return worklog

    def _create_jira_mappings(self):
        JiraUserMapping.objects.create(
            jira_account_id="acct-1",
            jira_display_name="Employee Jira",
            employee=self.employee.profile,
        )
        JiraProjectMapping.objects.create(
            jira_project_key="ALPHA",
            jira_project_name="Alpha Jira",
            project=self.project,
        )
        JiraIssueMapping.objects.create(
            jira_issue_key="ALPHA-1",
            jira_issue_id="50001",
            task=self.task,
        )

    def test_jira_settings_store_token_without_returning_secret(self):
        self._auth(self.manager)

        response = self.client.patch(
            reverse("core:time_jira_settings"),
            {
                "base_url": "https://example.atlassian.net/",
                "auth_email": "admin@example.com",
                "api_token": "secret-token",
                "enabled": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertNotIn("api_token", response.data)
        self.assertTrue(response.data["has_api_token"])
        connection = JiraConnection.objects.get()
        self.assertNotEqual(connection.api_token_encrypted, "secret-token")
        self.assertTrue(connection.api_token_encrypted.startswith("gAAAA"))
        self.assertEqual(connection.get_api_token(), "secret-token")

    def test_jira_mappings_endpoint_creates_grouped_mappings(self):
        self._auth(self.manager)

        user_response = self.client.post(
            reverse("core:time_jira_mappings"),
            {
                "mapping_type": "user",
                "jira_account_id": "acct-1",
                "jira_display_name": "Employee Jira",
                "employee_id": self.employee.profile.id,
            },
            format="json",
        )
        project_response = self.client.post(
            reverse("core:time_jira_mappings"),
            {
                "mapping_type": "project",
                "jira_project_key": "alpha",
                "jira_project_name": "Alpha Jira",
                "project_id": self.project.id,
            },
            format="json",
        )
        issue_response = self.client.post(
            reverse("core:time_jira_mappings"),
            {
                "mapping_type": "issue",
                "jira_issue_key": "alpha-1",
                "jira_issue_id": "50001",
                "task_id": self.task.id,
            },
            format="json",
        )
        list_response = self.client.get(reverse("core:time_jira_mappings"))

        self.assertEqual(user_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(project_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(issue_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(project_response.data["jira_project_key"], "ALPHA")
        self.assertEqual(issue_response.data["jira_issue_key"], "ALPHA-1")
        self.assertEqual(len(list_response.data["users"]), 1)
        self.assertEqual(len(list_response.data["projects"]), 1)
        self.assertEqual(len(list_response.data["issues"]), 1)

    def test_jira_preview_validates_missing_mappings(self):
        self._auth(self.manager)

        response = self.client.post(
            reverse("core:time_jira_import_preview"),
            {
                "date_from": "2026-05-18",
                "date_to": "2026-05-24",
                "worklogs": [self._jira_worklog()],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["error_count"], 1)
        codes = {
            message["code"]
            for message in response.data["rows"][0]["validation_messages"]
        }
        self.assertIn("missing_user_mapping", codes)
        self.assertIn("missing_project_mapping", codes)

    def test_jira_commit_imports_worklog_with_source_metadata(self):
        self._create_jira_mappings()
        self._auth(self.manager)

        response = self.client.post(
            reverse("core:time_jira_import_commit"),
            {
                "date_from": "2026-05-18",
                "date_to": "2026-05-24",
                "worklogs": [self._jira_worklog()],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["counts"]["created"], 1)
        entry = TimeEntry.objects.get(source_external_id="10001")
        self.assertEqual(entry.source_type, TimeEntrySourceType.JIRA)
        self.assertEqual(entry.employee, self.employee.profile)
        self.assertEqual(entry.project, self.project)
        self.assertEqual(entry.task, self.task)
        self.assertEqual(entry.hours, 2)
        self.assertEqual(entry.source_metadata["issue_key"], "ALPHA-1")
        self.assertEqual(entry.source_metadata["worklog_id"], "10001")
        batch = TimeImportBatch.objects.get(pk=response.data["batch_id"])
        self.assertEqual(batch.source_type, ImportBatchSource.JIRA)
        self.assertEqual(batch.committed_rows, 1)
        self.assertEqual(batch.rows.count(), 1)
        self.assertEqual(batch.rows.first().committed_entry, entry)

    def test_jira_commit_rerun_skips_unchanged_worklog(self):
        self._create_jira_mappings()
        self._auth(self.manager)
        payload = {
            "date_from": "2026-05-18",
            "date_to": "2026-05-24",
            "worklogs": [self._jira_worklog()],
        }

        first = self.client.post(
            reverse("core:time_jira_import_commit"),
            payload,
            format="json",
        )
        second = self.client.post(
            reverse("core:time_jira_import_commit"),
            payload,
            format="json",
        )

        self.assertEqual(first.status_code, status.HTTP_200_OK)
        self.assertEqual(second.status_code, status.HTTP_200_OK)
        self.assertEqual(second.data["counts"]["skipped"], 1)
        self.assertEqual(
            TimeEntry.objects.filter(source_external_id="10001").count(), 1
        )

    def _tempo_worklog(self, **overrides):
        worklog = {
            "tempoWorklogId": "tw-10001",
            "author": {
                "accountId": "tempo-user-1",
                "displayName": "Employee Tempo",
            },
            "account": {
                "id": "tempo-account-1",
                "key": "alpha-account",
                "name": "Alpha Account",
            },
            "project": {
                "id": "tempo-project-1",
                "key": "alpha",
                "name": "Alpha Tempo Project",
            },
            "team": {"id": "tempo-team-1", "name": "Alpha Team"},
            "issue": {"key": "ALPHA-1"},
            "startDate": "2026-05-18",
            "startTime": "09:30:00",
            "timeSpentSeconds": 7200,
            "description": "Tempo implementation",
            "updatedAt": "2026-05-18T11:00:00Z",
        }
        worklog.update(overrides)
        return worklog

    def _create_tempo_mappings(self):
        TempoUserMapping.objects.create(
            tempo_user_id="tempo-user-1",
            tempo_display_name="Employee Tempo",
            employee=self.employee.profile,
        )
        TempoAccountMapping.objects.create(
            tempo_account_id="tempo-account-1",
            tempo_account_key="ALPHA-ACCOUNT",
            tempo_account_name="Alpha Account",
            project=self.project,
        )
        TempoProjectMapping.objects.create(
            tempo_project_id="tempo-project-1",
            tempo_project_key="ALPHA",
            tempo_project_name="Alpha Tempo Project",
            project=self.project,
        )
        TempoTeamMapping.objects.create(
            tempo_team_id="tempo-team-1",
            tempo_team_name="Alpha Team",
            project=self.project,
        )

    def test_tempo_settings_store_token_without_returning_secret(self):
        self._auth(self.manager)

        response = self.client.patch(
            reverse("core:time_tempo_settings"),
            {
                "base_url": "https://api.tempo.io/4/",
                "api_token": "tempo-secret",
                "enabled": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertNotIn("api_token", response.data)
        self.assertTrue(response.data["has_api_token"])
        connection = TempoConnection.objects.get()
        self.assertNotEqual(connection.api_token_encrypted, "tempo-secret")
        self.assertTrue(connection.api_token_encrypted.startswith("gAAAA"))
        self.assertEqual(connection.get_api_token(), "tempo-secret")

    def test_tempo_test_connection_uses_posted_settings(self):
        self._auth(self.manager)
        tempo_response = Mock()
        tempo_response.status_code = 200
        tempo_response.json.return_value = {"results": []}

        with patch(
            "core.services.tempo_time_import_service.requests.get",
            return_value=tempo_response,
        ) as get:
            response = self.client.post(
                reverse("core:time_tempo_test_connection"),
                {
                    "base_url": "https://api.tempo.io/4/",
                    "api_token": "tempo-secret",
                    "enabled": False,
                },
                format="json",
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["last_test_status"], "success")
        self.assertTrue(response.data["has_api_token"])
        self.assertNotIn("api_token", response.data)
        connection = TempoConnection.objects.get()
        self.assertEqual(connection.base_url, "https://api.tempo.io/4")
        self.assertEqual(connection.get_api_token(), "tempo-secret")
        self.assertFalse(connection.enabled)
        self.assertEqual(
            get.call_args.kwargs["headers"]["Authorization"],
            "Bearer tempo-secret",
        )

    def test_tempo_mappings_endpoint_creates_grouped_mappings(self):
        self._auth(self.manager)

        user_response = self.client.post(
            reverse("core:time_tempo_mappings"),
            {
                "mapping_type": "user",
                "tempo_user_id": "tempo-user-1",
                "tempo_display_name": "Employee Tempo",
                "employee_id": self.employee.profile.id,
            },
            format="json",
        )
        account_response = self.client.post(
            reverse("core:time_tempo_mappings"),
            {
                "mapping_type": "account",
                "tempo_account_id": "tempo-account-1",
                "tempo_account_key": "alpha-account",
                "tempo_account_name": "Alpha Account",
                "project_id": self.project.id,
            },
            format="json",
        )
        project_response = self.client.post(
            reverse("core:time_tempo_mappings"),
            {
                "mapping_type": "project",
                "tempo_project_id": "tempo-project-1",
                "tempo_project_key": "alpha",
                "tempo_project_name": "Alpha Tempo Project",
                "project_id": self.project.id,
            },
            format="json",
        )
        team_response = self.client.post(
            reverse("core:time_tempo_mappings"),
            {
                "mapping_type": "team",
                "tempo_team_id": "tempo-team-1",
                "tempo_team_name": "Alpha Team",
                "project_id": self.project.id,
            },
            format="json",
        )
        list_response = self.client.get(reverse("core:time_tempo_mappings"))

        self.assertEqual(user_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(account_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(project_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(team_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(account_response.data["tempo_account_key"], "ALPHA-ACCOUNT")
        self.assertEqual(project_response.data["tempo_project_key"], "ALPHA")
        self.assertEqual(len(list_response.data["users"]), 1)
        self.assertEqual(len(list_response.data["accounts"]), 1)
        self.assertEqual(len(list_response.data["projects"]), 1)
        self.assertEqual(len(list_response.data["teams"]), 1)

    def test_tempo_project_discovery_uses_posted_token_and_suggests_project_id(self):
        self._auth(self.manager)

        def tempo_get(url, **kwargs):
            tempo_response = Mock()
            tempo_response.status_code = 200
            if url.endswith("/worklogs"):
                tempo_response.json.return_value = {
                    "results": [
                        {
                            "tempoWorklogId": 126,
                            "author": {"accountId": "tempo-user-1"},
                            "issue": {"id": 124},
                            "startDate": "2026-05-18",
                            "timeSpentSeconds": 7200,
                        }
                    ]
                }
            elif url.endswith("/accounts"):
                tempo_response.json.return_value = {
                    "metadata": {"next": None},
                    "results": [
                        {
                            "id": "tempo-account-1",
                            "key": "alpha-account",
                            "name": "Alpha Account",
                        }
                    ],
                }
            elif url.endswith("/projects"):
                tempo_response.json.return_value = {
                    "metadata": {"next": None},
                    "results": [
                        {
                            "id": "tempo-project-1",
                            "key": "alpha",
                            "name": "Alpha Tempo Project",
                        }
                    ],
                }
            elif url.endswith("/teams"):
                tempo_response.json.return_value = {
                    "metadata": {"next": None},
                    "results": [{"id": "tempo-team-1", "name": "Alpha Team"}],
                }
            return tempo_response

        with patch(
            "core.services.tempo_time_import_service.requests.get",
            side_effect=tempo_get,
        ) as get:
            response = self.client.post(
                reverse("core:time_tempo_project_discovery"),
                {
                    "base_url": "https://api.tempo.io/4/",
                    "api_token": "tempo-secret",
                    "date_from": "2026-05-18",
                    "date_to": "2026-05-24",
                },
                format="json",
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["counts"]["worklogs"], 1)
        self.assertEqual(
            get.call_args_list[0].kwargs["headers"]["Authorization"],
            "Bearer tempo-secret",
        )
        self.assertEqual(response.data["counts"]["accounts"], 1)
        self.assertEqual(response.data["counts"]["projects"], 1)
        self.assertEqual(response.data["counts"]["teams"], 1)
        self.assertEqual(
            response.data["projects"][0]["tempo_project_id"], "tempo-project-1"
        )
        self.assertEqual(response.data["projects"][0]["project_id"], self.project.id)
        self.assertEqual(
            response.data["projects"][0]["suggested_project"]["id"],
            self.project.id,
        )
        self.assertEqual(
            response.data["accounts"][0]["tempo_account_id"], "tempo-account-1"
        )
        self.assertFalse(TempoConnection.objects.exists())

    def test_tempo_project_discovery_returns_existing_mapping_project_id(self):
        self._auth(self.manager)
        connection = TempoConnection.get_solo()
        connection.set_api_token("saved-secret")
        connection.enabled = True
        connection.save()
        TempoAccountMapping.objects.create(
            tempo_account_id="tempo-account-1",
            tempo_account_key="ALPHA-ACCOUNT",
            tempo_account_name="Alpha Account",
            project=self.other_project,
        )

        def tempo_get(url, **kwargs):
            tempo_response = Mock()
            tempo_response.status_code = 200
            if url.endswith("/accounts"):
                tempo_response.json.return_value = {
                    "metadata": {"next": None},
                    "results": [
                        {
                            "id": "tempo-account-1",
                            "key": "ALPHA-ACCOUNT",
                            "name": "Alpha Account",
                        }
                    ],
                }
            elif url.endswith("/projects"):
                tempo_response.json.return_value = {
                    "metadata": {"next": None},
                    "results": [],
                }
            elif url.endswith("/teams"):
                tempo_response.json.return_value = {
                    "metadata": {"next": None},
                    "results": [],
                }
            else:
                tempo_response.json.return_value = {"results": []}
            return tempo_response

        with patch(
            "core.services.tempo_time_import_service.requests.get",
            side_effect=tempo_get,
        ) as get:
            response = self.client.post(
                reverse("core:time_tempo_project_discovery"),
                {
                    "date_from": "2026-05-18",
                    "date_to": "2026-05-24",
                },
                format="json",
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            get.call_args_list[0].kwargs["headers"]["Authorization"],
            "Bearer saved-secret",
        )
        self.assertEqual(
            response.data["accounts"][0]["project_id"], self.other_project.id
        )
        self.assertEqual(
            response.data["accounts"][0]["existing_mapping"]["project_id"],
            self.other_project.id,
        )

    def test_jira_project_discovery_uses_posted_token_and_suggests_mappings(self):
        self._auth(self.manager)

        def jira_get(url, **kwargs):
            jira_response = Mock()
            jira_response.status_code = 200
            if url.endswith("/rest/api/3/search/jql"):
                jira_response.json.return_value = {
                    "issues": [
                        {
                            "id": "50001",
                            "key": "ALPHA-1",
                            "fields": {
                                "summary": "Implementation",
                                "project": {"key": "ALPHA", "name": "Alpha"},
                            },
                        }
                    ],
                    "total": 1,
                }
            elif url.endswith("/rest/api/3/issue/ALPHA-1/worklog"):
                jira_response.json.return_value = {
                    "worklogs": [
                        {
                            "id": "10001",
                            "issueId": "50001",
                            "author": {
                                "accountId": "acct-1",
                                "displayName": "Employee",
                                "emailAddress": "employee@example.com",
                            },
                            "started": "2026-05-18T10:00:00.000+0000",
                            "timeSpentSeconds": 7200,
                        }
                    ],
                    "total": 1,
                }
            return jira_response

        with patch(
            "core.services.jira_time_import_service.requests.get",
            side_effect=jira_get,
        ) as get:
            response = self.client.post(
                reverse("core:time_jira_project_discovery"),
                {
                    "base_url": "https://bloomhub.atlassian.net",
                    "auth_email": "admin@example.com",
                    "api_token": "jira-secret",
                    "date_from": "2026-05-18",
                    "date_to": "2026-05-24",
                },
                format="json",
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["counts"]["worklogs"], 1)
        self.assertEqual(response.data["counts"]["users"], 1)
        self.assertEqual(response.data["counts"]["projects"], 1)
        self.assertEqual(response.data["counts"]["issues"], 1)
        self.assertEqual(get.call_args_list[0].kwargs["auth"][0], "admin@example.com")
        self.assertEqual(get.call_args_list[0].kwargs["auth"][1], "jira-secret")
        self.assertEqual(response.data["projects"][0]["jira_project_key"], "ALPHA")
        self.assertEqual(response.data["projects"][0]["project_id"], self.project.id)
        self.assertEqual(response.data["issues"][0]["jira_issue_key"], "ALPHA-1")
        self.assertEqual(response.data["issues"][0]["task_id"], self.task.id)
        self.assertEqual(response.data["users"][0]["jira_account_id"], "acct-1")
        self.assertEqual(
            response.data["users"][0]["employee_id"], self.employee.profile.id
        )
        self.assertFalse(JiraConnection.objects.exists())

    def test_jira_project_discovery_returns_existing_mapping_project_id(self):
        self._auth(self.manager)
        connection = JiraConnection.get_solo()
        connection.base_url = "https://saved.atlassian.net"
        connection.auth_email = "saved@example.com"
        connection.set_api_token("saved-secret")
        connection.enabled = True
        connection.save()
        JiraProjectMapping.objects.create(
            jira_project_key="ALPHA",
            jira_project_name="Alpha Jira",
            project=self.other_project,
        )

        def jira_get(url, **kwargs):
            jira_response = Mock()
            jira_response.status_code = 200
            if url.endswith("/rest/api/3/search/jql"):
                jira_response.json.return_value = {
                    "issues": [
                        {
                            "id": "50001",
                            "key": "ALPHA-1",
                            "fields": {
                                "summary": "Implementation",
                                "project": {"key": "ALPHA", "name": "Alpha Jira"},
                            },
                        }
                    ],
                    "total": 1,
                }
            else:
                jira_response.json.return_value = {"worklogs": [], "total": 0}
            return jira_response

        with patch(
            "core.services.jira_time_import_service.requests.get",
            side_effect=jira_get,
        ) as get:
            response = self.client.post(
                reverse("core:time_jira_project_discovery"),
                {
                    "date_from": "2026-05-18",
                    "date_to": "2026-05-24",
                },
                format="json",
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(get.call_args_list[0].kwargs["auth"][0], "saved@example.com")
        self.assertEqual(get.call_args_list[0].kwargs["auth"][1], "saved-secret")
        self.assertEqual(
            response.data["projects"][0]["project_id"], self.other_project.id
        )
        self.assertEqual(
            response.data["projects"][0]["existing_mapping"]["project_id"],
            self.other_project.id,
        )

    def test_tempo_preview_validates_missing_mappings(self):
        self._auth(self.manager)

        response = self.client.post(
            reverse("core:time_tempo_import_preview"),
            {
                "date_from": "2026-05-18",
                "date_to": "2026-05-24",
                "worklogs": [
                    self._tempo_worklog(
                        account={"id": "unmapped-account", "key": "unmapped"},
                        project={"id": "unmapped-project", "key": "unmapped"},
                        team={"id": "unmapped-team", "name": "Unmapped"},
                        issue={},
                    )
                ],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["error_count"], 1)
        codes = {
            message["code"]
            for message in response.data["rows"][0]["validation_messages"]
        }
        self.assertIn("missing_user_mapping", codes)
        self.assertIn("missing_project_mapping", codes)

    def test_tempo_preview_resolves_project_mapping_by_issue_project_key(self):
        self._auth(self.manager)
        TempoUserMapping.objects.create(
            tempo_user_id="tempo-user-1",
            tempo_display_name="Employee Tempo",
            employee=self.employee.profile,
        )
        TempoProjectMapping.objects.create(
            tempo_project_id="tempo-project-from-discovery",
            tempo_project_key="ALPHA",
            tempo_project_name="Alpha Tempo Project",
            project=self.project,
        )

        response = self.client.post(
            reverse("core:time_tempo_import_preview"),
            {
                "date_from": "2026-05-18",
                "date_to": "2026-05-24",
                "worklogs": [
                    self._tempo_worklog(
                        account={},
                        project={},
                        team={},
                        issue={"key": "ALPHA-2"},
                    )
                ],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["error_count"], 0)
        self.assertEqual(response.data["rows"][0]["project_id"], self.project.id)

    def test_tempo_preview_resolves_project_mapping_by_issue_key_in_comment(self):
        self._auth(self.manager)
        TempoUserMapping.objects.create(
            tempo_user_id="tempo-user-1",
            tempo_display_name="Employee Tempo",
            employee=self.employee.profile,
        )
        TempoProjectMapping.objects.create(
            tempo_project_id="tempo-project-from-discovery",
            tempo_project_key="BHB",
            tempo_project_name="BloomHub",
            project=self.project,
        )

        response = self.client.post(
            reverse("core:time_tempo_import_preview"),
            {
                "date_from": "2026-05-25",
                "date_to": "2026-05-31",
                "worklogs": [
                    {
                        "tempoWorklogId": "2",
                        "author": {"accountId": "tempo-user-1"},
                        "startDate": "2026-05-25",
                        "timeSpentSeconds": 18000,
                        "description": "Working on work item BHB-450",
                        "updatedAt": "2026-05-25T12:20:06Z",
                    }
                ],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["error_count"], 0)
        self.assertEqual(response.data["rows"][0]["jira_issue_key"], "BHB-450")
        self.assertEqual(response.data["rows"][0]["project_id"], self.project.id)
        self.assertEqual(response.data["rows"][0]["start_time"], None)
        self.assertEqual(response.data["rows"][0]["end_time"], None)

    def test_tempo_preview_resolves_project_from_jira_mapping_for_comment_issue_key(
        self,
    ):
        self._auth(self.manager)
        TempoUserMapping.objects.create(
            tempo_user_id="tempo-user-1",
            tempo_display_name="Employee Tempo",
            employee=self.employee.profile,
        )
        JiraProjectMapping.objects.create(
            jira_project_key="BHB",
            jira_project_name="BloomHub",
            project=self.project,
        )

        response = self.client.post(
            reverse("core:time_tempo_import_preview"),
            {
                "date_from": "2026-05-25",
                "date_to": "2026-05-31",
                "worklogs": [
                    {
                        "tempoWorklogId": "2",
                        "author": {"accountId": "tempo-user-1"},
                        "startDate": "2026-05-25",
                        "timeSpentSeconds": 18000,
                        "description": "Working on work item BHB-450",
                    }
                ],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["error_count"], 0)
        self.assertEqual(response.data["rows"][0]["project_id"], self.project.id)

    def test_tempo_commit_imports_worklog_with_source_metadata(self):
        self._create_tempo_mappings()
        self._auth(self.manager)

        response = self.client.post(
            reverse("core:time_tempo_import_commit"),
            {
                "date_from": "2026-05-18",
                "date_to": "2026-05-24",
                "worklogs": [self._tempo_worklog()],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["counts"]["created"], 1)
        entry = TimeEntry.objects.get(source_external_id="tw-10001")
        self.assertEqual(entry.source_type, TimeEntrySourceType.TEMPO)
        self.assertEqual(entry.employee, self.employee.profile)
        self.assertEqual(entry.project, self.project)
        self.assertEqual(entry.task, self.task)
        self.assertEqual(entry.hours, 2)
        self.assertEqual(entry.start_time, time(9, 30))
        self.assertEqual(entry.end_time, time(11, 30))
        self.assertEqual(entry.source_metadata["tempo_worklog_id"], "tw-10001")
        self.assertEqual(entry.source_metadata["tempo_account_id"], "tempo-account-1")
        self.assertEqual(entry.source_metadata["jira_issue_key"], "ALPHA-1")
        self.assertEqual(entry.source_metadata["start_time"], "09:30:00")
        self.assertEqual(entry.source_metadata["end_time"], "11:30:00")
        batch = TimeImportBatch.objects.get(pk=response.data["batch_id"])
        self.assertEqual(batch.source_type, ImportBatchSource.TEMPO)
        self.assertEqual(batch.committed_rows, 1)
        self.assertEqual(batch.rows.count(), 1)
        self.assertEqual(batch.rows.first().committed_entry, entry)

    def test_tempo_commit_rerun_skips_unchanged_worklog(self):
        self._create_tempo_mappings()
        self._auth(self.manager)
        payload = {
            "date_from": "2026-05-18",
            "date_to": "2026-05-24",
            "worklogs": [self._tempo_worklog()],
        }

        first = self.client.post(
            reverse("core:time_tempo_import_commit"),
            payload,
            format="json",
        )
        second = self.client.post(
            reverse("core:time_tempo_import_commit"),
            payload,
            format="json",
        )

        self.assertEqual(first.status_code, status.HTTP_200_OK)
        self.assertEqual(second.status_code, status.HTTP_200_OK)
        self.assertEqual(second.data["counts"]["skipped"], 1)
        self.assertEqual(
            TimeEntry.objects.filter(source_external_id="tw-10001").count(),
            1,
        )

    def test_tempo_duplicate_against_manual_entry_is_skipped(self):
        self._create_tempo_mappings()
        TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.project,
            task=self.task,
            work_date=date(2026, 5, 18),
            hours="2.00",
            notes="Tempo implementation",
            duplicate_fingerprint="",
        )
        manual = TimeEntry.objects.get(notes="Tempo implementation")
        from core.services.time_tracking_service import fingerprint_for_entry

        manual.duplicate_fingerprint = fingerprint_for_entry(manual)
        manual.save(update_fields=["duplicate_fingerprint"])
        self._auth(self.manager)

        response = self.client.post(
            reverse("core:time_tempo_import_commit"),
            {
                "date_from": "2026-05-18",
                "date_to": "2026-05-24",
                "worklogs": [self._tempo_worklog()],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["counts"]["skipped"], 1)
        self.assertFalse(
            TimeEntry.objects.filter(
                source_type=TimeEntrySourceType.TEMPO,
                source_external_id="tw-10001",
            ).exists()
        )

    def _csv_file(self, name="timesheet.csv", rows=None):
        if rows is None:
            rows = [
                "employee,date,project,jira issue,hours,notes",
                "employee,2026-05-18,Alpha,ALPHA-1,2.50,Document work",
            ]
        content = "\n".join(rows).encode("utf-8")
        return SimpleUploadedFile(name, content, content_type="text/csv")

    def test_document_upload_rejects_empty_and_unsupported_files(self):
        self._auth(self.manager)

        empty_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {"file": SimpleUploadedFile("empty.csv", b"", content_type="text/csv")},
            format="multipart",
        )
        unsupported_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {
                "file": SimpleUploadedFile(
                    "timesheet.txt",
                    b"employee,date,hours\nEmployee,2026-05-18,2",
                    content_type="text/plain",
                )
            },
            format="multipart",
        )

        self.assertEqual(empty_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(unsupported_response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_document_csv_upload_preview_and_commit(self):
        self._auth(self.manager)

        upload_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {"file": self._csv_file()},
            format="multipart",
        )
        self.assertEqual(upload_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(upload_response.data["source_type"], "document_import")
        self.assertEqual(upload_response.data["valid_rows"], 1)

        commit_response = self.client.post(
            reverse("core:time_import_batch_commit", args=[upload_response.data["id"]]),
            format="json",
        )

        self.assertEqual(commit_response.status_code, status.HTTP_200_OK)
        self.assertEqual(commit_response.data["committed_rows"], 1)
        entry = TimeEntry.objects.get(source_type=TimeEntrySourceType.DOCUMENT_IMPORT)
        self.assertEqual(entry.employee, self.employee.profile)
        self.assertEqual(entry.project, self.project)
        self.assertEqual(entry.task, self.task)
        self.assertEqual(entry.source_metadata["file_name"], "timesheet.csv")
        self.assertEqual(entry.source_metadata["row_number"], 2)

    def test_document_revalidation_preserves_committed_rows(self):
        from core.services.document_time_import_service import validate_batch_rows

        self._auth(self.manager)
        upload_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {"file": self._csv_file()},
            format="multipart",
        )
        commit_response = self.client.post(
            reverse("core:time_import_batch_commit", args=[upload_response.data["id"]]),
            format="json",
        )
        batch = TimeImportBatch.objects.get(pk=upload_response.data["id"])
        row = batch.rows.get()
        committed_entry = row.committed_entry

        validate_batch_rows(batch)
        batch.refresh_from_db()
        row.refresh_from_db()

        self.assertEqual(commit_response.status_code, status.HTTP_200_OK)
        self.assertEqual(row.status, ImportRowStatus.COMMITTED)
        self.assertEqual(row.committed_entry, committed_entry)
        self.assertEqual(batch.committed_rows, 1)
        self.assertEqual(batch.skipped_rows, 0)

    def test_document_column_mapping_can_fix_ambiguous_headers(self):
        self._auth(self.manager)
        upload_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {
                "file": self._csv_file(
                    rows=[
                        "person,day,client,duration,description",
                        "employee,2026-05-18,Alpha,2.00,Mapped work",
                    ]
                )
            },
            format="multipart",
        )
        self.assertEqual(upload_response.status_code, status.HTTP_201_CREATED)

        map_response = self.client.post(
            reverse(
                "core:time_document_import_map_columns",
                args=[upload_response.data["id"]],
            ),
            {
                "column_mapping": {
                    "employee": "person",
                    "date": "day",
                    "project": "client",
                    "hours": "duration",
                    "notes": "description",
                }
            },
            format="json",
        )

        self.assertEqual(map_response.status_code, status.HTTP_200_OK)
        self.assertEqual(map_response.data["valid_rows"], 1)
        self.assertEqual(map_response.data["error_rows"], 0)

    def test_document_import_accepts_exported_id_columns(self):
        self._auth(self.manager)
        upload_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {
                "file": self._csv_file(
                    rows=[
                        "employee,employee_id,date,start_time,end_time,project,project_id,task,task_id,jira_issue_key,hours,notes,source_type,status",
                        f"Johnas Doe,{self.employee.profile.id},2026-05-25,09:15:00,14:15:00,Atlas,{self.project.id},,,BHB-450,5.00,Working on work item BHB-450,tempo,approved",
                    ]
                )
            },
            format="multipart",
        )

        self.assertEqual(upload_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(upload_response.data["valid_rows"], 1)
        self.assertEqual(upload_response.data["error_rows"], 0)
        batch = TimeImportBatch.objects.get(pk=upload_response.data["id"])
        row = batch.rows.get()
        self.assertEqual(row.status, ImportRowStatus.VALID)
        self.assertEqual(row.parsed_data["employee_id"], self.employee.profile.id)
        self.assertEqual(row.parsed_data["project_id"], self.project.id)
        self.assertEqual(row.parsed_data["jira_issue_key"], "BHB-450")
        self.assertEqual(row.parsed_data["start_time"], "09:15:00")
        self.assertEqual(row.parsed_data["end_time"], "14:15:00")

    def test_document_import_commit_retains_exported_start_and_end_time(self):
        self._auth(self.manager)
        upload_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {
                "file": self._csv_file(
                    rows=[
                        "employee,employee_id,date,start_time,end_time,project,project_id,task,task_id,jira_issue_key,hours,notes",
                        f"Johnas Doe,{self.employee.profile.id},2026-05-25,09:15:00,14:15:00,Atlas,{self.project.id},,,BHB-450,5.00,Roundtrip row",
                    ]
                )
            },
            format="multipart",
        )
        commit_response = self.client.post(
            reverse("core:time_import_batch_commit", args=[upload_response.data["id"]]),
            format="json",
        )

        self.assertEqual(commit_response.status_code, status.HTTP_200_OK)
        entry = TimeEntry.objects.get(source_type=TimeEntrySourceType.DOCUMENT_IMPORT)
        self.assertEqual(entry.start_time, time(9, 15))
        self.assertEqual(entry.end_time, time(14, 15))
        self.assertEqual(entry.source_metadata["start_time"], "09:15:00")
        self.assertEqual(entry.source_metadata["end_time"], "14:15:00")

    def test_document_import_partial_commit_skips_duplicate_and_keeps_errors(self):
        self._auth(self.manager)
        duplicate = TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.project,
            task=self.task,
            work_date=date(2026, 5, 18),
            hours="2.00",
            notes="Duplicate row",
            duplicate_fingerprint="",
        )
        from core.services.time_tracking_service import fingerprint_for_entry

        duplicate.duplicate_fingerprint = fingerprint_for_entry(duplicate)
        duplicate.save(update_fields=["duplicate_fingerprint"])

        upload_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {
                "file": self._csv_file(
                    rows=[
                        "employee,date,project,jira issue,hours,notes",
                        "employee,2026-05-18,Alpha,ALPHA-1,2.00,Duplicate row",
                        "unknown,2026-05-19,Alpha,ALPHA-1,3.00,Bad user",
                        "employee,2026-05-20,Alpha,ALPHA-1,4.00,Valid row",
                    ]
                )
            },
            format="multipart",
        )
        commit_response = self.client.post(
            reverse("core:time_import_batch_commit", args=[upload_response.data["id"]]),
            format="json",
        )

        self.assertEqual(commit_response.status_code, status.HTTP_200_OK)
        self.assertEqual(commit_response.data["committed_rows"], 1)
        self.assertEqual(commit_response.data["skipped_rows"], 1)
        self.assertEqual(commit_response.data["error_rows"], 1)
        self.assertEqual(commit_response.data["status"], "partially_committed")
        self.assertEqual(
            TimeEntry.objects.filter(
                source_type=TimeEntrySourceType.DOCUMENT_IMPORT
            ).count(),
            1,
        )

    def test_document_commit_ignores_duplicates_and_commits_valid_rows(self):
        self._auth(self.manager)
        duplicate = TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.project,
            task=self.task,
            work_date=date(2026, 5, 18),
            hours="2.00",
            notes="Duplicate row",
            duplicate_fingerprint="",
        )
        from core.services.time_tracking_service import fingerprint_for_entry

        duplicate.duplicate_fingerprint = fingerprint_for_entry(duplicate)
        duplicate.save(update_fields=["duplicate_fingerprint"])

        upload_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {
                "file": self._csv_file(
                    rows=[
                        "employee,date,project,jira issue,hours,notes",
                        "employee,2026-05-18,Alpha,ALPHA-1,2.00,Duplicate row",
                        "employee,2026-05-20,Alpha,ALPHA-1,4.00,Valid row",
                    ]
                )
            },
            format="multipart",
        )
        commit_response = self.client.post(
            reverse("core:time_import_batch_commit", args=[upload_response.data["id"]]),
            format="json",
        )

        self.assertEqual(commit_response.status_code, status.HTTP_200_OK)
        self.assertEqual(commit_response.data["committed_rows"], 1)
        self.assertEqual(commit_response.data["skipped_rows"], 1)
        self.assertEqual(commit_response.data["error_rows"], 0)
        self.assertEqual(commit_response.data["status"], "partially_committed")
        self.assertEqual(
            TimeEntry.objects.filter(
                source_type=TimeEntrySourceType.DOCUMENT_IMPORT,
                notes="Valid row",
            ).count(),
            1,
        )

    def test_document_commit_succeeds_when_all_rows_are_duplicates(self):
        self._auth(self.manager)
        duplicate = TimeEntry.objects.create(
            employee=self.employee.profile,
            project=self.project,
            task=self.task,
            work_date=date(2026, 5, 18),
            hours="2.00",
            notes="Duplicate row",
            duplicate_fingerprint="",
        )
        from core.services.time_tracking_service import fingerprint_for_entry

        duplicate.duplicate_fingerprint = fingerprint_for_entry(duplicate)
        duplicate.save(update_fields=["duplicate_fingerprint"])

        upload_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {
                "file": self._csv_file(
                    rows=[
                        "employee,date,project,jira issue,hours,notes",
                        "employee,2026-05-18,Alpha,ALPHA-1,2.00,Duplicate row",
                    ]
                )
            },
            format="multipart",
        )
        commit_response = self.client.post(
            reverse("core:time_import_batch_commit", args=[upload_response.data["id"]]),
            format="json",
        )

        self.assertEqual(commit_response.status_code, status.HTTP_200_OK)
        self.assertEqual(commit_response.data["committed_rows"], 0)
        self.assertEqual(commit_response.data["skipped_rows"], 1)
        self.assertEqual(commit_response.data["error_rows"], 0)
        self.assertEqual(commit_response.data["status"], "committed")

    def test_document_xlsx_and_docx_upload_parse_tables(self):
        import openpyxl
        from docx import Document as DocxDocument

        self._auth(self.manager)

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["employee", "date", "project", "jira issue", "hours", "notes"])
        sheet.append(["employee", "2026-05-18", "Alpha", "ALPHA-1", "1.50", "XLSX"])
        xlsx_buffer = io.BytesIO()
        workbook.save(xlsx_buffer)
        xlsx_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {
                "file": SimpleUploadedFile(
                    "timesheet.xlsx",
                    xlsx_buffer.getvalue(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                )
            },
            format="multipart",
        )

        document = DocxDocument()
        table = document.add_table(rows=2, cols=6)
        for index, value in enumerate(
            ["employee", "date", "project", "jira issue", "hours", "notes"]
        ):
            table.rows[0].cells[index].text = value
        for index, value in enumerate(
            ["employee", "2026-05-19", "Alpha", "ALPHA-1", "1.25", "DOCX"]
        ):
            table.rows[1].cells[index].text = value
        docx_buffer = io.BytesIO()
        document.save(docx_buffer)
        docx_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {
                "file": SimpleUploadedFile(
                    "timesheet.docx",
                    docx_buffer.getvalue(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "wordprocessingml.document"
                    ),
                )
            },
            format="multipart",
        )

        self.assertEqual(xlsx_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(docx_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(TimeImportBatch.objects.count(), 2)

    def test_weekly_dashboard_includes_all_time_entry_sources(self):
        self._time_entry(notes="Manual", hours="1.00")
        self._time_entry(
            notes="Jira",
            hours="2.00",
            source_type=TimeEntrySourceType.JIRA,
            source_external_id="jira-dashboard",
            source_metadata={"jira_issue_key": "ALPHA-1"},
        )
        self._time_entry(
            notes="Tempo",
            hours="3.00",
            source_type=TimeEntrySourceType.TEMPO,
            source_external_id="tempo-dashboard",
            source_metadata={"jira_issue_key": "ALPHA-1"},
        )
        self._time_entry(
            notes="Document",
            hours="4.00",
            source_type=TimeEntrySourceType.DOCUMENT_IMPORT,
            source_external_id="document-dashboard",
            source_metadata={"jira_issue_key": "ALPHA-1"},
        )
        self._auth(self.manager)

        response = self.client.get(
            reverse("core:time_tracking_weekly_dashboard"),
            {"week_start": "2026-05-18"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["total_hours"], "10.00")
        self.assertEqual(
            response.data["totals_by_source"][TimeEntrySourceType.MANUAL], "1.00"
        )
        self.assertEqual(
            response.data["totals_by_source"][TimeEntrySourceType.JIRA], "2.00"
        )
        self.assertEqual(
            response.data["totals_by_source"][TimeEntrySourceType.TEMPO], "3.00"
        )
        self.assertEqual(
            response.data["totals_by_source"][TimeEntrySourceType.DOCUMENT_IMPORT],
            "4.00",
        )

    def test_approval_queue_filters_and_scopes_to_manager_team(self):
        self._time_entry(
            status=TimeEntryStatus.SUBMITTED,
            notes="Needs approval",
            source_type=TimeEntrySourceType.JIRA,
            source_external_id="jira-approval",
            source_metadata={"jira_issue_key": "ALPHA-1"},
        )
        other_user = self._make_user("otheremployee", self.employee_role)
        self._time_entry(
            employee=other_user.profile,
            status=TimeEntryStatus.SUBMITTED,
            notes="Not direct report",
            source_type=TimeEntrySourceType.TEMPO,
            source_external_id="tempo-other-approval",
            source_metadata={"jira_issue_key": "ALPHA-1"},
        )
        self._auth(self.manager)

        response = self.client.get(
            reverse("core:time_tracking_approval_queue"),
            {"source_type": TimeEntrySourceType.JIRA},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]["notes"], "Needs approval")

    def test_export_timesheets_csv_and_xlsx_permissions(self):
        self._time_entry(
            status=TimeEntryStatus.APPROVED,
            start_time=time(9, 15),
            notes="Export me",
            source_type=TimeEntrySourceType.DOCUMENT_IMPORT,
            source_external_id="document-export",
            source_metadata={"jira_issue_key": "ALPHA-1", "worklog_id": "doc-row"},
        )
        self._auth(self.outsider)
        forbidden = self.client.get(
            reverse("core:time_tracking_timesheet_export"),
            {"format": "csv", "date_from": "2026-05-18", "date_to": "2026-05-18"},
        )

        self._auth(self.manager)
        csv_response = self.client.get(
            reverse("core:time_tracking_timesheet_export"),
            {"format": "csv", "date_from": "2026-05-18", "date_to": "2026-05-18"},
        )
        xlsx_response = self.client.get(
            reverse("core:time_tracking_timesheet_export"),
            {"format": "xlsx", "date_from": "2026-05-18", "date_to": "2026-05-18"},
        )

        self.assertEqual(forbidden.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(csv_response.status_code, status.HTTP_200_OK)
        csv_content = csv_response.content.decode("utf-8")
        self.assertIn("Export me", csv_content)
        self.assertIn("source_external_id", csv_content)
        self.assertIn("start_time", csv_content)
        self.assertIn("09:15:00", csv_content)
        self.assertEqual(xlsx_response.status_code, status.HTTP_200_OK)
        self.assertIn(
            "spreadsheetml.sheet",
            xlsx_response["Content-Type"],
        )

    def test_planned_vs_actual_returns_allocation_variance(self):
        ProjectAssignment.objects.create(
            user_profile=self.employee.profile,
            project=self.project,
            allocation_percentage=50,
            start_date=date(2026, 5, 18),
            status=ProjectAssignmentStatus.ACTIVE,
        )
        self._time_entry(hours="12.00", notes="Actual")
        self._auth(self.manager)

        response = self.client.get(
            reverse("core:time_tracking_planned_vs_actual"),
            {
                "date_from": "2026-05-18",
                "date_to": "2026-05-24",
                "employee_id": self.employee.profile.id,
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        row = response.data["rows"][0]
        self.assertEqual(row["planned_hours"], "20.00")
        self.assertEqual(row["actual_hours"], "12.00")
        self.assertEqual(row["variance_hours"], "-8.00")

    def test_planned_vs_actual_includes_allocated_employee_without_time(self):
        ProjectAssignment.objects.create(
            user_profile=self.employee.profile,
            project=self.project,
            allocation_percentage=50,
            start_date=date(2026, 5, 18),
            status=ProjectAssignmentStatus.ACTIVE,
        )
        self._auth(self.manager)

        response = self.client.get(
            reverse("core:time_tracking_planned_vs_actual"),
            {
                "date_from": "2026-05-18",
                "date_to": "2026-05-24",
                "employee_id": self.employee.profile.id,
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        row = response.data["rows"][0]
        self.assertEqual(row["employee_id"], self.employee.profile.id)
        self.assertEqual(row["planned_hours"], "20.00")
        self.assertEqual(row["actual_hours"], "0.00")
        self.assertEqual(row["variance_hours"], "-20.00")

    def test_import_batch_list_and_detail_filters_for_admins(self):
        self._auth(self.manager)
        upload_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {"file": self._csv_file()},
            format="multipart",
        )
        batch_id = upload_response.data["id"]

        list_response = self.client.get(
            reverse("core:time_import_batch_list"),
            {"source_type": TimeEntrySourceType.DOCUMENT_IMPORT},
        )
        detail_response = self.client.get(
            reverse("core:time_import_batch_detail", args=[batch_id])
        )

        self.assertEqual(list_response.status_code, status.HTTP_200_OK)
        self.assertEqual(detail_response.status_code, status.HTTP_200_OK)
        self.assertEqual(list_response.data[0]["id"], batch_id)
        self.assertEqual(detail_response.data["file_name"], "timesheet.csv")

    def test_import_row_validation_codes_are_consistent(self):
        self._auth(self.manager)
        upload_response = self.client.post(
            reverse("core:time_document_import_upload"),
            {
                "file": self._csv_file(
                    rows=[
                        "employee,date,project,hours,notes",
                        "unknown,bad-date,Missing,abc,Bad row",
                    ]
                )
            },
            format="multipart",
        )

        self.assertEqual(upload_response.status_code, status.HTTP_201_CREATED)
        messages = upload_response.data["rows"][0]["validation_messages"]
        codes = {message["code"] for message in messages}
        self.assertEqual(
            codes,
            {"missing_user", "missing_project", "invalid_date", "invalid_hours"},
        )

    def test_source_change_review_accept_current_creates_audit_event(self):
        entry = self._time_entry(
            source_type=TimeEntrySourceType.JIRA,
            source_external_id="jira-review",
            source_metadata={
                "jira_issue_key": "ALPHA-1",
                "source_change_flag": TimeEntrySourceChangeFlag.REVIEW_REQUIRED,
            },
        )
        self._auth(self.manager)

        queue_response = self.client.get(
            reverse("core:time_tracking_source_change_review")
        )
        resolve_response = self.client.post(
            reverse("core:time_tracking_source_change_resolve", args=[entry.id]),
            {"action": "accept_current", "note": "Reviewed"},
            format="json",
        )

        self.assertEqual(queue_response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(queue_response.data), 1)
        self.assertEqual(resolve_response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            resolve_response.data["source_metadata"]["source_change_flag"],
            TimeEntrySourceChangeFlag.NONE,
        )
        self.assertTrue(
            TimeEntryAuditEvent.objects.filter(
                time_entry=entry,
                event_type=TimeEntryAuditEventType.SOURCE_CHANGED,
                metadata__action="accept_current",
            ).exists()
        )

    def test_source_change_review_applies_pending_update_for_unapproved_row(self):
        entry = self._time_entry(
            source_type=TimeEntrySourceType.TEMPO,
            source_external_id="tempo-review",
            source_metadata={
                "jira_issue_key": "ALPHA-1",
                "source_change_flag": TimeEntrySourceChangeFlag.REVIEW_REQUIRED,
                "source_pending_update": {
                    "work_date": "2026-05-19",
                    "hours": "3.50",
                    "notes": "Updated from source",
                },
            },
        )
        self._auth(self.manager)

        response = self.client.post(
            reverse("core:time_tracking_source_change_resolve", args=[entry.id]),
            {"action": "apply_source"},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        entry.refresh_from_db()
        self.assertEqual(entry.work_date, date(2026, 5, 19))
        self.assertEqual(str(entry.hours), "3.50")
        self.assertEqual(entry.notes, "Updated from source")
        self.assertEqual(
            entry.source_metadata["source_change_flag"],
            TimeEntrySourceChangeFlag.NONE,
        )

    def test_external_api_rate_limit_errors_are_admin_visible(self):
        from rest_framework.exceptions import ValidationError

        from core.services.jira_time_import_service import (
            JiraImportFilters,
            fetch_jira_worklogs,
        )

        connection = JiraConnection.get_solo()
        connection.base_url = "https://example.atlassian.net"
        connection.auth_email = "admin@example.com"
        connection.set_api_token("secret")
        connection.enabled = True
        connection.save()
        response = Mock()
        response.status_code = 429
        response.headers = {"Retry-After": "1"}

        with (
            patch("core.services.jira_time_import_service.sleep"),
            patch(
                "core.services.jira_time_import_service.requests.get",
                return_value=response,
            ),
        ):
            with self.assertRaises(ValidationError) as raised:
                fetch_jira_worklogs(
                    connection,
                    JiraImportFilters(
                        date_from=date(2026, 5, 18),
                        date_to=date(2026, 5, 24),
                    ),
                )

        self.assertIn("Jira rate-limit", str(raised.exception))

    def test_jira_fetch_pages_search_and_worklog_results(self):
        from core.services.jira_time_import_service import (
            JiraImportFilters,
            fetch_jira_worklogs,
        )

        connection = JiraConnection.get_solo()
        connection.base_url = "https://example.atlassian.net"
        connection.auth_email = "admin@example.com"
        connection.set_api_token("secret")
        connection.enabled = True
        connection.save()

        def response(payload):
            mock_response = Mock()
            mock_response.status_code = 200
            mock_response.headers = {}
            mock_response.json.return_value = payload
            return mock_response

        def fake_get(url, **kwargs):
            params = kwargs.get("params") or {}
            start_at = params.get("startAt", 0)
            if url.endswith("/worklog/deleted") or url.endswith("/worklog/updated"):
                return response({"values": []})
            if url.endswith("/rest/api/3/search/jql"):
                if params.get("nextPageToken") == "page-2":
                    return response(
                        {
                            "issues": [{"key": "ALPHA-2", "id": "50002"}],
                            "isLast": True,
                        }
                    )
                return response(
                    {
                        "issues": [{"key": "ALPHA-1", "id": "50001"}],
                        "nextPageToken": "page-2",
                        "isLast": False,
                    }
                )
            if url.endswith("/issue/ALPHA-1/worklog"):
                worklogs = [
                    self._jira_worklog(id="10001"),
                    self._jira_worklog(id="10002"),
                ]
                return response(
                    {"worklogs": worklogs[start_at : start_at + 1], "total": 2}
                )
            return response({"worklogs": [self._jira_worklog(id="20001")], "total": 1})

        with patch("core.services.jira_time_import_service.requests.get", fake_get):
            worklogs = fetch_jira_worklogs(
                connection,
                JiraImportFilters(
                    date_from=date(2026, 5, 18),
                    date_to=date(2026, 5, 24),
                ),
            )

        self.assertEqual(
            [(worklog["issueKey"], worklog["id"]) for worklog in worklogs],
            [("ALPHA-1", "10001"), ("ALPHA-1", "10002"), ("ALPHA-2", "20001")],
        )

    def test_jira_fetch_skips_deleted_worklogs_without_existing_entry(self):
        from core.services.jira_time_import_service import (
            JiraImportFilters,
            fetch_jira_worklogs,
        )

        connection = JiraConnection.get_solo()
        connection.base_url = "https://example.atlassian.net"
        connection.auth_email = "admin@example.com"
        connection.set_api_token("secret")
        connection.enabled = True
        connection.save()

        def response(payload):
            mock_response = Mock()
            mock_response.status_code = 200
            mock_response.headers = {}
            mock_response.json.return_value = payload
            return mock_response

        def fake_get(url, **kwargs):
            if url.endswith("/worklog/deleted"):
                return response(
                    {
                        "values": [
                            {"worklogId": "10000"},
                            {"worklogId": "10001"},
                        ]
                    }
                )
            if url.endswith("/worklog/updated"):
                return response({"values": []})
            if url.endswith("/rest/api/3/search/jql"):
                return response({"issues": [], "isLast": True})
            return response({"worklogs": [], "total": 0})

        self._time_entry(
            source_type=TimeEntrySourceType.JIRA,
            source_external_id="10001",
            source_metadata={
                "jira_issue_key": "ALPHA-1",
                "issue_id": "50001",
                "worklog_id": "10001",
                "author_account_id": "acct-1",
                "author_display_name": "Employee Jira",
                "started": "2026-05-18T09:00:00.000+0000",
                "time_spent_seconds": 7200,
                "comment": "Jira implementation",
                "updated": "2026-05-18T11:00:00.000+0000",
            },
        )

        with patch("core.services.jira_time_import_service.requests.get", fake_get):
            worklogs = fetch_jira_worklogs(
                connection,
                JiraImportFilters(
                    date_from=date(2026, 5, 18),
                    date_to=date(2026, 5, 24),
                ),
            )

        self.assertEqual(len(worklogs), 1)
        self.assertEqual(worklogs[0]["id"], "10001")
        self.assertEqual(worklogs[0]["issueKey"], "ALPHA-1")
        self.assertTrue(worklogs[0]["deleted"])

    def test_tempo_timeout_errors_are_admin_visible(self):
        from rest_framework.exceptions import ValidationError

        from core.services.tempo_time_import_service import (
            TempoImportFilters,
            fetch_tempo_worklogs,
        )

        connection = TempoConnection.get_solo()
        connection.base_url = "https://api.tempo.io/4"
        connection.set_api_token("secret")
        connection.enabled = True
        connection.save()

        import requests

        with (
            patch("core.services.tempo_time_import_service.sleep"),
            patch(
                "core.services.tempo_time_import_service.requests.get",
                side_effect=requests.Timeout("network down"),
            ),
        ):
            with self.assertRaises(ValidationError) as raised:
                fetch_tempo_worklogs(
                    connection,
                    TempoImportFilters(
                        date_from=date(2026, 5, 18),
                        date_to=date(2026, 5, 24),
                    ),
                )

        self.assertIn("network down", str(raised.exception))
